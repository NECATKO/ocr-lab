"""Word (.docx) dışa aktarıcı: kitabın sayfasına benzer çıktı (spec §2 Hedef 2, onaylanan spike `spike_docx.py`).

Her taranmış sayfa bir Word bölümü; satırlar elle satır sonuyla (Shift+Enter) kırılır; kitap sayfa numarası bölümün
altlığında. Harfler değişmez: yalnız satır sonuna tire eklenir ve dipnot numarası üst simge yapılır.
"""
import math
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from ocrnow.export_txt import FAILED_TEXT, page_label, plain_text, split_blocks
from ocrnow.layout import Block, Page

BOLD_DEFAULT = r"\[\d+[ab] [^\]]*\]"   # yazma nüsha işaretleri: [106b sayfanın sağ yanı]

FONT = "Times New Roman"
BODY_PT = 10.5            # en büyük gövde puntosu
FOOTNOTE_PT = 9.0
HEADER_PT = 9.5           # Page-header ve altlıktaki numara
HEADER_MIN_PT = 6.0       # bundan uzun başlık küçülmez, Word kaydırır
HEADER_SPACING = 30       # Page-header harf aralığı (twip)
SEPARATOR_PT = 8.0
TEXT_WIDTH_PT = 391       # 17 cm - 2 × 1.6 cm kenar boşluğu
CHAR_EM = 0.42            # Times New Roman ortalama harf genişliği (punto başına)
INDENT = Cm(0.9)          # paragrafın ilk satır girintisi
WHOLE_BOLD = frozenset({"Title", "Section-header"})
Ranges = list[tuple[int, int]]


def _half_down(pt: float) -> float:
    """0.5'e aşağı yuvarlar (yukarı yuvarlamak satırı taşırıp Word'ün kendiliğinden kaydırmasına yol açabilirdi)."""
    return math.floor(pt * 2) / 2


def _fit(max_chars: int, cap: float, indent_pt: float = 0.0) -> float:
    """max_chars karakterlik satırın (girinti payı indent_pt ile) metin genişliğine sığdığı punto, en çok cap."""
    if max_chars <= 0:
        return cap
    return min(cap, _half_down((TEXT_WIDTH_PT - indent_pt) / (CHAR_EM * max_chars)))


def font_size_for(max_chars: int) -> float:
    return _fit(max_chars, BODY_PT)


def _spans(block: Block, text: str) -> list[tuple[int, int, bool]]:
    """Bloğun satırları (başlangıç, bitiş, tire eklenecek mi); satır yapısı yoksa bütün metin tek satır."""
    if block.lines and block.category != "Table":
        return [(line.start, line.end, line.hyphen) for line in block.lines]
    return [(0, len(text), False)]


def _size_for(blocks: list[Block], cap: float, indented: bool) -> float:
    """Blokların bütün satırlarının sığdığı punto (R18). Satır uzunluğuna eklenen tire, girintili ilk satıra girinti
    payı girer. Satır yapısı olan satırlar sayılır: 2+ satırlı blokların satırları ve cap'te sığan tek satırlar.
    cap'te sığmayan tek satırlı blok (R11 geri dönüşü, Table, lines == [] olanlar) akan paragraftır: Word kaydırır."""
    size = cap
    for b in blocks:
        indent_pt = INDENT.pt if indented and b.indent else 0.0
        fits = [_fit(e - s + h, cap, indent_pt if k == 0 else 0.0)
                for k, (s, e, h) in enumerate(_spans(b, plain_text(b)))]
        if len(fits) > 1 or fits[0] == cap:
            size = min(size, *fits)
    return size


def _header_size(block: Block) -> float:
    """Page-header puntosu: en uzun satırı harf aralığıyla sığmazsa yalnız başlık küçülür (en az HEADER_MIN_PT)."""
    longest = max(e - s + h for s, e, h in _spans(block, block.text))
    if longest <= 0:
        return HEADER_PT
    spacing_pt = HEADER_SPACING / 20
    return max(HEADER_MIN_PT, min(HEADER_PT, _half_down((TEXT_WIDTH_PT / longest - spacing_pt) / CHAR_EM)))


def _covers(ranges: Ranges, pos: int) -> bool:
    return any(a <= pos < b for a, b in ranges)


def _spans_break(ranges: Ranges, end: int) -> bool:
    """Aralık satır sonunu aşıyor mu (kırılımdan önceki ve sonraki harfi birlikte kapsıyor mu): eklenen tire de
    o aralığa aittir (bölünmüş kelime ya da nüsha işareti)."""
    return any(a < end < b for a, b in ranges)


def _sup_pattern(page: Page) -> re.Pattern | None:
    """Sayfanın dipnot numaraları; gövdede bir harfin hemen ardından geldiklerinde (taştında724) üst simge olurlar."""
    nums = {m.group(1) for b in page.blocks if b.category == "Footnote" and (m := re.match(r"(\d+)\s", b.text))}
    if not nums:
        return None
    return re.compile(r"(?<=[^\W\d_])(?:" + "|".join(sorted(nums, key=len, reverse=True)) + r")(?!\d)")


def _add_runs(p, block: Block, page_sup: re.Pattern | None, bold_re: re.Pattern | None, highlight: bool,
              size: float | None = None, spacing: int | None = None) -> None:
    """Bloğu satır satır run'lara yazar: run'lar kalın, üst simge ve sarı aralıklarının sınırlarından bölünür."""
    text = plain_text(block)
    if block.category in WHOLE_BOLD:
        bold: Ranges = [(0, len(text))]
    else:
        bold = [m.span() for m in bold_re.finditer(text)] if bold_re else []
    sup = [m.span() for m in page_sup.finditer(text)] if page_sup else []
    if block.category == "Footnote" and (m := re.match(r"\d+(?=\s)", text)):
        sup.append(m.span())
    marks = [(m.start, m.end) for m in block.marks] if highlight and block.category != "Table" else []
    spans = _spans(block, text)
    for k, (start, end, hyphen) in enumerate(spans):
        cuts = sorted({start, end, *(x for r in bold + sup + marks for x in r if start < x < end)})
        parts = [(text[a:b], (_covers(bold, a), _covers(sup, a), _covers(marks, a))) for a, b in zip(cuts, cuts[1:])]
        if hyphen:
            parts.append(("-", (_spans_break(bold, end), False, _spans_break(marks, end))))
        merged: list[list] = []
        for s, props in parts:
            if merged and merged[-1][1] == props:
                merged[-1][0] += s
            else:
                merged.append([s, props])
        runs = []
        for s, (is_bold, is_sup, is_yellow) in merged or [["", (False, False, False)]]:
            r = p.add_run(s)
            if spacing:     # harf aralığı; rPr'de ilk eklenir, python-docx sonrakileri şema sırasına yerleştirir
                sp = OxmlElement("w:spacing")
                sp.set(qn("w:val"), str(spacing))
                r._r.get_or_add_rPr().append(sp)
            if is_bold:
                r.bold = True
            if is_sup:
                r.font.superscript = True
            if is_yellow:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            if size:
                r.font.size = Pt(size)
            runs.append(r)
        if k < len(spans) - 1:
            runs[-1].add_break(WD_BREAK.LINE)


def _setup_section(section) -> None:
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width, section.page_height = Cm(17), Cm(24)
    section.left_margin = section.right_margin = Cm(1.6)
    section.top_margin, section.bottom_margin = Cm(1.5), Cm(1.8)
    section.footer_distance = Cm(0.8)


def _footer(section, page: Page) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if page.book_page_side == "sag" else WD_ALIGN_PARAGRAPH.LEFT
    if label := page_label(page):
        p.add_run(label).font.size = Pt(HEADER_PT)


def _header(doc, block: Block):
    """Page-header: ortalı, küçük punto, harf aralığı 30; işaretleme yapılmaz."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    _add_runs(p, block, None, None, False, _header_size(block), HEADER_SPACING)
    return p


def _close_section(doc, last_paragraph) -> None:
    """Bölümü son paragrafında kapatır. doc.add_section() araya ayrı bir boş paragraf koyar; dolu bir sayfada o
    paragraf fazladan boş bir Word sayfası açabilirdi."""
    sentinel = doc.element.body.get_or_add_sectPr()
    last_paragraph._p.set_sectPr(sentinel.clone())
    for ref in sentinel.xpath("w:headerReference|w:footerReference"):
        sentinel.remove(ref)            # sonraki bölüm kendi altlığını alır


def write_docx(pages: list[Page], out: Path, highlight: bool = True, bold_pattern: str | None = BOLD_DEFAULT) -> Path:
    groups = [split_blocks(page) for page in pages]
    body_pt = _size_for([b for _, body, _ in groups for b in body], BODY_PT, indented=True)
    footnote_pt = min(body_pt, _size_for([b for _, _, fns in groups for b in fns], FOOTNOTE_PT, indented=False))
    bold_re = re.compile(bold_pattern) if bold_pattern else None

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name, style.font.size = FONT, Pt(body_pt)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.15
    for n, (page, (headers, body, footnotes)) in enumerate(zip(pages, groups)):
        headers, body, footnotes = ([b for b in g if plain_text(b)] for g in (headers, body, footnotes))
        section = doc.sections[-1]
        _setup_section(section)
        _footer(section, page)
        sup = _sup_pattern(page)
        last = None
        for block in headers:
            last = _header(doc, block)
        for block in body:
            last = doc.add_paragraph()
            last.paragraph_format.space_after = Pt(5)
            if block.indent:
                last.paragraph_format.first_line_indent = INDENT
            _add_runs(last, block, sup, bold_re, highlight)
        if footnotes:
            sep = doc.add_paragraph()
            sep.paragraph_format.space_before = Pt(8)
            sep.add_run("_" * 18).font.size = Pt(SEPARATOR_PT)
            for block in footnotes:
                last = doc.add_paragraph()
                _add_runs(last, block, sup, bold_re, highlight, footnote_pt)
        if page.failed:
            last = doc.add_paragraph(FAILED_TEXT)
        if last is None:
            last = doc.add_paragraph()      # bloksuz (ya da yalnız numaralı) sayfa: boş paragraf
        if n < len(pages) - 1:
            _close_section(doc, last)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out
