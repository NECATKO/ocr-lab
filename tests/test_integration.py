"""Uçtan uca (M11): ocr-now CLI → gerçek collect_jobs/render/process_job/önbellek → gerçek docx/txt dışa aktarıcıları.

Yalnız model sunucusu (engines.server) ve tesseract sahtedir; GPU ve tesseract gerekmez.
"""
import json
import sys
import types
from contextlib import contextmanager

import pypdfium2 as pdfium
import pytest
from docx import Document

from ocrnow import engines, paths
from ocrnow.cli import main


def make_pdf(path, n):
    pdf = pdfium.PdfDocument.new()
    for _ in range(n):
        pdf.new_page(200, 300)
    pdf.save(str(path))
    pdf.close()
    return path


@pytest.fixture
def world(tmp_path, monkeypatch):
    for name, sub in [("INPUT", "input"), ("CACHE", "output/_cache"), ("PAGES", "output/_pages"),
                      ("DOCS_OUT", "output/belgeler"), ("LOGS", "output/_logs"), ("GT", "gt"),
                      ("SETTINGS", "ayarlar")]:
        monkeypatch.setattr(paths, name, tmp_path / sub)
    monkeypatch.setattr(engines, "missing", lambda name: None)
    monkeypatch.setattr(sys.stdin, "isatty", lambda: False)
    state = types.SimpleNamespace(footers={}, fail=set(), starts=0, start_error=None, tsv_calls=0)

    def tsv(img):                     # satır yok: bloklar akan paragraf
        state.tsv_calls += 1
        return ""

    monkeypatch.setattr(engines, "tesseract_tsv", tsv)

    @contextmanager
    def server(name):
        state.starts += 1
        if state.start_error:
            raise state.start_error

        def run(img, mode):
            n = int(img.stem.removeprefix("syf"))                   # render: .../syf<N>.png
            if n in state.fail:
                raise RuntimeError("model yanıt vermedi")
            blocks = [{"category": "Text", "bbox": [50, 50, 500, 400], "text": f"metin {n}"}]
            if state.footers.get(n) is not None:
                blocks.append({"category": "Page-footer", "bbox": [50, 760, 120, 790], "text": str(state.footers[n])})
            return json.dumps(blocks)
        yield run

    monkeypatch.setattr(engines, "server", server)
    state.pdf = make_pdf(tmp_path / "k.pdf", 3)
    return state


def footers(docx_path):
    return [s.footer.paragraphs[0].text for s in Document(docx_path).sections]


def test_numbering_gap_inferred_end_to_end(world, capsys):
    world.footers = {1: 412, 3: 414}
    assert main([str(world.pdf), "-f", "docx", "txt"]) == 0
    assert footers(paths.DOCS_OUT / "k.docx") == ["412", "[413]", "414"]
    txt = (paths.DOCS_OUT / "k.txt").read_text(encoding="utf-8")
    assert txt == "--- s. 412 ---\nmetin 1\n\n--- s. [413] ---\nmetin 2\n\n--- s. 414 ---\nmetin 3\n"
    assert "tahmin edilen sayfa numaraları:\n  [413]" in capsys.readouterr().out
    assert world.tsv_calls == 3                                # gerçek tesseract çağrılmadı


def test_failed_page_has_placeholder_end_to_end(world, capsys):
    world.footers, world.fail = {1: 412, 2: 413, 3: 414}, {2}
    assert main([str(world.pdf), "-f", "docx", "txt"]) == 1
    paras = [p.text for p in Document(paths.DOCS_OUT / "k.docx").paragraphs]
    assert paras == ["metin 1", "[Bu sayfa okunamadı]", "metin 3"]
    assert footers(paths.DOCS_OUT / "k.docx") == ["412", "[413]", "414"]
    assert "--- s. [413] ---\n[Bu sayfa okunamadı]" in (paths.DOCS_OUT / "k.txt").read_text(encoding="utf-8")
    assert "s. 2: okunamadı: model yanıt vermedi" in capsys.readouterr().out


def test_all_pages_failed_writes_nothing_end_to_end(world, capsys):
    world.fail = {1, 2, 3}
    assert main([str(world.pdf), "-f", "docx", "txt"]) == 1
    assert not paths.DOCS_OUT.exists() or list(paths.DOCS_OUT.iterdir()) == []
    assert "k: hiçbir sayfa okunamadı, belge yazılmadı" in capsys.readouterr().out


def test_server_start_failure_end_to_end(world, capsys):
    world.start_error = RuntimeError("sunucu ayağa kalkmadı")
    assert main([str(world.pdf), "-f", "txt"]) == 1
    err = capsys.readouterr().err
    assert "'dots.mocr' modeli başlatılamadı: sunucu ayağa kalkmadı" in err and err.count("Hata:") == 1
    assert world.starts == 1 and not paths.DOCS_OUT.exists()


def test_all_cached_rerun_end_to_end(world):
    world.footers = {1: 412, 2: 413, 3: 414}
    assert main([str(world.pdf), "-p", "2-3", "-f", "docx"]) == 0
    first = paths.DOCS_OUT / "k_s2-3.docx"
    assert footers(first) == ["413", "414"] and world.starts == 1
    world.start_error = AssertionError("hepsi önbellekteyken sunucu açılmamalı")
    assert main([str(world.pdf), "-p", "2-3", "-f", "docx", "txt"]) == 0
    assert world.starts == 1                                   # model hiç çağrılmadı
    again = paths.DOCS_OUT / "k_s2-3_2.docx"                    # eski belge ezilmedi
    assert [p.text for p in Document(again).paragraphs] == [p.text for p in Document(first).paragraphs]
    assert (paths.DOCS_OUT / "k_s2-3_2.txt").exists()
