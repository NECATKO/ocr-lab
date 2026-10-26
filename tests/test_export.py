from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Pt

from ocrnow import paths
from ocrnow.export_docx import write_docx, font_size_for
from ocrnow.export_txt import write_txt
from ocrnow.layout import Block, Line, Mark, Page
from ocrnow.pages import SourcePage
from ocrnow.pipeline import Options, process_page

SPIKE_SIZE = (1653, 2339)


def page(n, side="sol", inferred=False, text="bir iki üç dört", lines=None, marks=()):
    b = Block("Text", (0, 0, 10, 10), text, lines or [Line(0, 7, True), Line(8, 15, False)], True, list(marks))
    return Page("k.pdf", n, 100, 100, "dots.mocr", "duzen", [b], book_page=str(400 + n),
                book_page_inferred=inferred, book_page_side=side)


def blocks_page(*blocks, book_page="5"):
    return Page("k.pdf", 1, 100, 100, "dots.mocr", "duzen", list(blocks), book_page=book_page)


def bold_text(d):
    return "".join(r.text for p in d.paragraphs for r in p.runs if r.bold).replace("\n", "")


def test_docx_sections_and_footers(tmp_path):
    d = Document(write_docx([page(1), page(2, "sag", True)], tmp_path / "a.docx"))
    assert len(d.sections) == 2
    assert d.sections[0].footer.paragraphs[0].text == "401"
    assert d.sections[1].footer.paragraphs[0].text == "[402]"
    assert d.sections[0].footer.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert d.sections[1].footer.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert not d.sections[1].footer.is_linked_to_previous


def test_docx_page_setup(tmp_path):
    d = Document(write_docx([page(1), page(2)], tmp_path / "a.docx"))
    for s in d.sections:
        assert (round(s.page_width.cm, 2), round(s.page_height.cm, 2)) == (17, 24)
        assert round(s.left_margin.cm, 2) == round(s.right_margin.cm, 2) == 1.6
    assert d.styles["Normal"].font.name == "Times New Roman"


def test_no_empty_paragraph_at_section_break(tmp_path):
    """Bölüm sonu ayrı boş paragraf olmaz (dolu sayfada fazladan boş Word sayfası açardı)."""
    d = Document(write_docx([page(1), page(2)], tmp_path / "a.docx"))
    assert [p.text for p in d.paragraphs] == ["bir iki-\nüç dört"] * 2


def test_docx_line_breaks_and_added_hyphen(tmp_path):
    d = Document(write_docx([page(1)], tmp_path / "a.docx"))
    assert d.paragraphs[0].text == "bir iki-\nüç dört"
    assert round(d.paragraphs[0].paragraph_format.first_line_indent.cm, 2) == 0.9


def test_block_without_lines_is_one_line(tmp_path):
    p = blocks_page(Block("Text", None, "akan paragraf metni"))
    d = Document(write_docx([p], tmp_path / "a.docx"))
    assert d.paragraphs[0].text == "akan paragraf metni" and d.paragraphs[0].paragraph_format.first_line_indent is None


def test_highlight_toggle(tmp_path):
    p = page(1, marks=[Mark(4, 7, "belirsiz")])
    on = Document(write_docx([p], tmp_path / "on.docx"))
    off = Document(write_docx([p], tmp_path / "off.docx", highlight=False))
    yellow = lambda d: [r.text for r in d.paragraphs[0].runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW]
    assert yellow(on) == ["iki"] and yellow(off) == []


def test_overlapping_marks_highlight_union(tmp_path):
    p = page(1, text="mes‘ele var", lines=[Line(0, 11, False)],
             marks=[Mark(0, 7, "belirsiz"), Mark(3, 4, "kesme")])
    d = Document(write_docx([p], tmp_path / "a.docx"))
    runs = d.paragraphs[0].runs
    assert "".join(r.text for r in runs) == "mes‘ele var"
    assert [r.text for r in runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW] == ["mes‘ele"]


def test_highlight_spanning_hyphenated_break(tmp_path):
    """Satır sonunda bölünen sarı kelimenin eklenen tiresi de sarıdır."""
    p = page(1, text="hazâkat var", lines=[Line(0, 4, True), Line(4, 11, False)], marks=[Mark(0, 7, "belirsiz")])
    d = Document(write_docx([p], tmp_path / "a.docx"))
    yellow = [r.text for r in d.paragraphs[0].runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW]
    assert "".join(yellow).replace("\n", "") == "hazâ-kat"


def test_bold_marker_split_across_lines(tmp_path):
    t = "x [106b sayfanın sağ yanı] y"   # metinde birleşik kelime; satır "sayfa" + eklenen tire ile biter
    p = page(1, text=t, lines=[Line(0, 13, True), Line(13, len(t), False)])
    d = Document(write_docx([p], tmp_path / "a.docx"))
    assert d.paragraphs[0].text == "x [106b sayfa-\nnın sağ yanı] y"
    assert bold_text(d) == "[106b sayfa-nın sağ yanı]"
    none = Document(write_docx([p], tmp_path / "b.docx", bold_pattern=None))
    assert not any(r.bold for r in none.paragraphs[0].runs)


def test_bold_marker_split_after_kept_hyphen(tmp_path):
    t = "x [106b sayfa-nın sağ yanı] y"   # ana model tireyi korumuş
    p = page(1, text=t, lines=[Line(0, 14, False), Line(14, len(t), False)])
    d = Document(write_docx([p], tmp_path / "a.docx"))
    assert d.paragraphs[0].text == "x [106b sayfa-\nnın sağ yanı] y"
    assert bold_text(d) == "[106b sayfa-nın sağ yanı]"


def test_footnote_separator_and_superscripts(tmp_path):
    body = "taştında724 ve 724 sayı"
    p = blocks_page(Block("Footnote", None, "724 açıklama metni"),
                    Block("Text", None, body, [Line(0, len(body), False)]))
    d = Document(write_docx([p], tmp_path / "a.docx"))
    texts = [q.text for q in d.paragraphs]
    assert texts == [body, "_" * 18, "724 açıklama metni"]      # dipnot sayfa sonunda, harfler değişmez
    sup = [r.text for r in d.paragraphs[0].runs if r.font.superscript]
    assert sup == ["724"]                                     # yalnız harften hemen sonra gelen
    fn = d.paragraphs[2].runs
    assert fn[0].text == "724" and fn[0].font.superscript and fn[1].font.superscript is not True
    assert all(r.font.size == Pt(9) for r in fn)


def test_header_title_table(tmp_path):
    p = blocks_page(Block("Page-header", None, "KİTAP ADI"),
                    Block("Title", None, "BAŞLIK"),
                    Block("Table", None, "<table><tr><td>a</td><td>b &amp; c</td></tr></table>"),
                    Block("Page-footer", None, "5"))
    d = Document(write_docx([p], tmp_path / "a.docx"))
    head, title, table = d.paragraphs
    assert head.text == "KİTAP ADI" and head.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert head.runs[0].font.size == Pt(9.5)
    assert head.runs[0]._r.rPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}spacing") \
        .get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") == "30"
    assert title.text == "BAŞLIK" and all(r.bold for r in title.runs)
    assert table.text == "a b & c"


def test_blank_page_still_section(tmp_path):
    p = Page("k.pdf", 1, 100, 100, "dots.mocr", "duzen", [], book_page="7")
    assert Document(write_docx([p], tmp_path / "b.docx")).sections[0].footer.paragraphs[0].text == "7"


def test_failed_and_unnumbered_pages(tmp_path):
    failed = Page("k.pdf", 2, 0, 0, "dots.mocr", "duzen", [], book_page="PDF s. 2", book_page_inferred=True)
    none = Page("k.pdf", 3, 100, 100, "dots.mocr", "duzen", [Block("Page-footer", None, "x")])
    d = Document(write_docx([page(1), failed, none], tmp_path / "sub" / "b.docx"))
    assert len(d.sections) == 3
    assert [s.footer.paragraphs[0].text for s in d.sections] == ["401", "[PDF s. 2]", ""]


def test_failed_page_has_placeholder(tmp_path):
    """R22: okunamayan sayfa belgede boş sayfa değil, açık bir yer tutucudur."""
    failed = Page("k.pdf", 2, 0, 0, "dots.mocr", "duzen", [], book_page="402", book_page_inferred=True, failed=True)
    d = Document(write_docx([page(1), failed, page(3)], tmp_path / "a.docx"))
    assert [p.text for p in d.paragraphs] == ["bir iki-\nüç dört", "[Bu sayfa okunamadı]", "bir iki-\nüç dört"]
    assert [s.footer.paragraphs[0].text for s in d.sections] == ["401", "[402]", "403"]
    txt = write_txt([page(1), failed], tmp_path / "a.txt").read_text(encoding="utf-8")
    assert txt.endswith("--- s. [402] ---\n[Bu sayfa okunamadı]\n")


def test_long_lines_shrink_font():
    assert font_size_for(85) == 10.5 and font_size_for(120) <= 8
    assert font_size_for(100) * 2 == int(font_size_for(100) * 2)       # 0.5'lik adım
    assert font_size_for(100) * 0.42 * 100 <= 391                       # en uzun satır sığar


def pts(paragraph) -> set[float]:
    """Paragrafın run puntoları (run'da punto yoksa Normal stilinki)."""
    normal = paragraph.part.document.styles["Normal"].font.size.pt
    return {r.font.size.pt if r.font.size else normal for r in paragraph.runs}


def test_docx_font_is_per_section(tmp_path):
    """R23: punto bölüm (sayfa) başına; bir sayfadaki uzun satır öteki sayfaları küçültmez."""
    long = "b " + "a" * 100                          # girintisiz ikinci satır, eklenen tireyle 100 karakter
    p2 = page(2, text=long, lines=[Line(0, 1, False), Line(2, 101, True), Line(101, 102, False)])
    d = Document(write_docx([page(1), p2, page(3)], tmp_path / "a.docx"))
    assert [pts(p) for p in d.paragraphs] == [{10.5}, {font_size_for(100)}, {10.5}]
    assert font_size_for(100) == 9.0


def test_font_floor_and_warning(tmp_path):
    """R23: punto 8'in altına inmez; inecek satır varsa sayfa/blok uyarısı verilir."""
    long = "b " + "a" * 160
    p2 = page(2, text=long, lines=[Line(0, 1, False), Line(2, 162, False)])
    p2.blocks.insert(0, Block("Page-header", None, "BAŞLIK"))
    warnings: list[str] = []
    d = Document(write_docx([page(1), p2], tmp_path / "a.docx", warnings=warnings))
    assert pts(d.paragraphs[0]) == {10.5} and pts(d.paragraphs[2]) == {8.0}
    assert warnings == ["docx sayfa 402, blok 2: satır 8 pt'de de sığmıyor (160 karakter), Word kaydırabilir"]
    assert write_docx([page(1)], tmp_path / "b.docx") == tmp_path / "b.docx"   # uyarı listesi isteğe bağlı


def test_txt(tmp_path):
    txt = write_txt([page(1), page(2, inferred=True)], tmp_path / "a.txt").read_text()
    assert txt.startswith("--- s. 401 ---\nbir iki üç dört") and "--- s. [402] ---" in txt


def test_txt_layout(tmp_path):
    p = blocks_page(Block("Page-header", None, "KİTAP"),
                    Block("Footnote", None, "1 dipnot"),
                    Block("Text", None, "birinci", [Line(0, 3, True), Line(3, 7, False)], marks=[Mark(0, 3, "belirsiz")]),
                    Block("Table", None, "<table><tr><td>a</td></tr></table>"),
                    Block("Picture", None, "[Görsel]"),
                    Block("Page-footer", None, "5"))
    out = write_txt([p, Page("k.pdf", 2, 0, 0, "dots.mocr", "duzen", [], book_page="6")], tmp_path / "x" / "a.txt")
    assert out == tmp_path / "x" / "a.txt"
    assert out.read_text(encoding="utf-8") == "--- s. 5 ---\nbirinci\n\na\n\n[Görsel]\n\n1 dipnot\n\n--- s. 6 ---\n"


def test_spike_page_docx(spike_layout, spike_tsv, tmp_path, monkeypatch):
    """Görev 10 Adım 5 (R7): spike sayfası docx'e; output/_spike/gorev10.docx kullanıcıya açması için kalır."""
    monkeypatch.setattr(paths, "CACHE", tmp_path / "cache")
    monkeypatch.setattr(paths, "GT", tmp_path / "gt")            # gerçek kelime listesi sızmaz: belirlenimci
    monkeypatch.setattr(paths, "SETTINGS", tmp_path / "ayarlar")
    p = process_page(Path("x.png"), SourcePage(Path("k.pdf"), 1), Options(), lambda img, m: spike_layout,
                     lambda img: spike_tsv, SPIKE_SIZE)
    out = write_docx([p], paths.OUTPUT / "_spike" / "gorev10.docx")
    d = Document(out)
    assert len(d.sections) == 1
    footer = d.sections[0].footer.paragraphs[0]
    assert footer.text == "412" and footer.alignment == WD_ALIGN_PARAGRAPH.LEFT
    paras = d.paragraphs
    assert paras[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    sep = [q.text for q in paras].index("_" * 18)
    body = paras[1:sep]
    assert sum(q.text.count("\n") + 1 for q in body) == 36
    runs = [r for q in paras for r in q.runs]
    assert any(r.bold and r.text.startswith("[106b") for r in runs)
    assert any(r.font.highlight_color == WD_COLOR_INDEX.YELLOW for r in runs)
    assert paras[sep + 1].runs[0].font.superscript
    longest = max(len(line) for q in body for line in q.text.split("\n"))
    assert set().union(*(pts(q) for q in body)) == {max(8.0, font_size_for(longest))}


def body_pt(d) -> float:
    """İlk gövde paragrafının (tek) puntosu."""
    (pt,) = pts(next(p for p in d.paragraphs if p.alignment != WD_ALIGN_PARAGRAPH.CENTER))
    return pt


def test_flowing_blocks_do_not_shrink_font(tmp_path):
    """R18: satır yapısı olmayan uzun tek-span bloklar akan paragraftır (Word kaydırır), punto hesabına girmez."""
    fallback = "a " * 210                                                       # R11 geri dönüşü: tek satır
    p = blocks_page(Block("Text", None, fallback.strip(), [Line(0, len(fallback.strip()), False)]),
                    Block("Table", None, "<table><tr><td>" + "b " * 100 + "</td></tr></table>",
                          [Line(0, 10, False), Line(11, 20, False)]),
                    Block("Caption", None, "c" * 120),
                    Block("Footnote", None, "1 " + "d " * 100))
    d = Document(write_docx([page(1), p], tmp_path / "a.docx"))
    assert all(pts(q) == {10.5} for q in d.paragraphs[:3])
    assert d.paragraphs[-1].runs[-1].font.size == Pt(9)                         # dipnot puntosu da küçülmez


def test_indent_counts_in_first_line_width(tmp_path):
    t = "f" * 86 + " g"
    lines = [Line(0, 86, False), Line(87, 88, False)]
    b = Block("Text", None, t, lines, True)
    indented = Document(write_docx([blocks_page(b)], tmp_path / "b.docx"))
    assert body_pt(indented) == 10.0
    b.indent = False
    assert body_pt(Document(write_docx([blocks_page(b)], tmp_path / "c.docx"))) == 10.5


def test_long_header_shrinks_only_header(tmp_path):
    p = blocks_page(Block("Page-header", None, "H" * 80), Block("Text", None, "kısa"))
    d = Document(write_docx([p], tmp_path / "a.docx"))
    assert d.paragraphs[0].runs[0].font.size == Pt(8) and body_pt(d) == 10.5
