"""SPIKE (throwaway): dots.mocr layout text + tesseract line structure -> docx (option A)."""
import csv, json, re, sys
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

S = Path(sys.argv[1])
blocks = json.loads((S / "dots_layout_syf1.json").read_text())

# --- tesseract lines: bbox + text
words = {}
lines = {}
for r in csv.DictReader(open(S / "tess_syf1.tsv"), delimiter="\t", quoting=csv.QUOTE_NONE):
    key = (r["block_num"], r["par_num"], r["line_num"])
    if r["level"] == "4":
        x, y, w, h = (int(r[k]) for k in ("left", "top", "width", "height"))
        lines[key] = dict(box=(x, y, x + w, y + h), words=[])
    elif r["level"] == "5" and r["text"].strip():
        lines[key]["words"].append(r["text"])
tlines = [dict(box=l["box"], text=" ".join(l["words"])) for l in lines.values() if l["words"]]


def inside(lb, bb, pad=15):
    cx, cy = (lb[0] + lb[2]) / 2, (lb[1] + lb[3]) / 2
    return bb[0] - pad <= cx <= bb[2] + pad and bb[1] - pad <= cy <= bb[3] + pad


def align_map(a, b):
    """Levenshtein hizalaması: a'daki her konum için b'deki karşılık gelen konum."""
    n, m = len(a), len(b)
    D = [list(range(m + 1))] + [[i] + [0] * m for i in range(1, n + 1)]
    for i in range(1, n + 1):
        ai, Di, Dp = a[i - 1], D[i], D[i - 1]
        for j in range(1, m + 1):
            Di[j] = min(Dp[j] + 1, Di[j - 1] + 1, Dp[j - 1] + (ai != b[j - 1]))
    amap = [0] * (n + 1)
    i, j = n, m
    amap[n] = m
    while i > 0:
        if j > 0 and D[i][j] == D[i - 1][j - 1] + (a[i - 1] != b[j - 1]):
            i, j = i - 1, j - 1
        elif D[i][j] == D[i - 1][j] + 1:
            i -= 1
        else:
            j -= 1
            continue
        amap[i] = j
    return amap


HYPH = "-.¬,"  # tesseract satır sonu tiresini bazen . ya da , okuyor


def rebreak(dtext, tl):
    """dots metnini tesseract satırlarına göre böl. Döner: satır listesi."""
    T, breaks = "", []
    for k, l in enumerate(tl):
        t = l["text"]
        hyph = k < len(tl) - 1 and len(t) > 1 and t[-1] in HYPH and t[-2].isalpha() and tl[k + 1]["text"][:1].islower()
        if hyph:
            t = t[:-1]
        T += t
        if k < len(tl) - 1:
            breaks.append((len(T), hyph))
            T += "" if hyph else " "
    D = " ".join(dtext.split())
    amap = align_map(T, D)
    out, cuts = [], []
    for pos, hyph in breaks:
        j = amap[pos]
        if hyph and (j >= len(D) or " " in D[max(0, j - 2):j + 2]):
            hyph = False                          # dots'ta burada kelime sınırı var: tire değil (ör. "ki,")
        if hyph:
            if j > 0 and D[j - 1] == "-":
                cuts.append((j, j, ""))          # dots tireyi korumuş: tireden sonra kır
            elif j < len(D) and D[j] == "-":
                cuts.append((j + 1, j + 1, ""))
            else:
                cuts.append((j, j, "-"))         # birleşik kelimeyi böl, tire ekle
        else:
            # en yakın boşluğu bul
            cand = [k for k in range(max(0, j - 4), min(len(D), j + 5)) if D[k] == " "]
            k = min(cand, key=lambda k: abs(k - j)) if cand else j
            cuts.append((k, k + 1 if cand else k, ""))
    prev = 0
    for a, b, suffix in cuts:
        out.append(D[prev:a] + suffix)
        prev = b
    out.append(D[prev:])
    return out


# --- docx
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(17), Cm(24)
sec.left_margin = sec.right_margin = Cm(1.6)
sec.top_margin, sec.bottom_margin = Cm(1.5), Cm(1.8)
sec.footer_distance = Cm(0.8)
st = doc.styles["Normal"]
st.font.name, st.font.size = "Times New Roman", Pt(10.5)
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.line_spacing = 1.15

fn_nums = [m.group(1) for b in blocks if b["category"] == "Footnote" for m in [re.match(r"(\d+)\s", b["text"])] if m]
MARK = re.compile(r"(\[\d+[ab] [^\]]*\])")                          # yazma nüsha işaretleri (kalın)
SUP = re.compile(r"(?<=[^\W\d])(" + "|".join(fn_nums) + r")(?=\s|$)") if fn_nums else None


def add_runs(p, text, size=None):
    # kalın işaretler satır sonunda bölünebilir: "[106b sayfa-" / "nın sağ yanı]"
    for part in re.split(r"(\[[^\]]*$|^[^\[]*\]|\[\d+[ab] [^\]]*\])", text):
        if not part:
            continue
        bold = part.startswith("[") or part.endswith("]")
        subs = SUP.split(part) if SUP else [part]
        for s in subs:
            if not s:
                continue
            r = p.add_run(s)
            r.bold = bold
            if SUP and s in fn_nums:
                r.font.superscript = True
            if size:
                r.font.size = Pt(size)


report = []
for b in blocks:
    cat, txt = b["category"], b.get("text", "")
    tl = [l for l in tlines if inside(l["box"], b["bbox"])]
    if cat == "Page-header":
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(txt)
        r.font.size = Pt(9.5)
        rpr = r._r.get_or_add_rPr(); sp = OxmlElement("w:spacing"); sp.set(qn("w:val"), "30"); rpr.append(sp)
    elif cat == "Text":
        rows = rebreak(txt, tl) if len(tl) > 1 else [txt]
        report.append((b["bbox"], len(tl), rows))
        indent = tl and tl[0]["box"][0] - b["bbox"][0] > 30
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        if indent:
            p.paragraph_format.first_line_indent = Cm(0.9)
        for k, row in enumerate(rows):
            add_runs(p, row)
            if k < len(rows) - 1:
                p.runs[-1].add_break(WD_BREAK.LINE)
    elif cat == "Footnote":
        sepp = doc.add_paragraph()
        sepp.paragraph_format.space_before = Pt(8)
        sepp.add_run("_" * 18).font.size = Pt(8)
        p = doc.add_paragraph()
        m = re.match(r"(\d+)\s+(.*)", txt)
        if m:
            r = p.add_run(m.group(1)); r.font.superscript = True; r.font.size = Pt(9)
            p.add_run("  " + m.group(2)).font.size = Pt(9)
        else:
            p.add_run(txt).font.size = Pt(9)
    elif cat == "Page-footer":
        # sayfa numarası altlığa: her taranmış sayfa kendi bölümü, kendi altlığı (çok sayfada: doc.add_section + is_linked_to_previous=False)
        ft = doc.sections[-1].footer
        ft.is_linked_to_previous = False
        fp = ft.paragraphs[0]
        left = b["bbox"][0] < 700  # kitapta numara solda mı sağda mı
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT if left else WD_ALIGN_PARAGRAPH.RIGHT
        fp.add_run(txt).font.size = Pt(9.5)

doc.save(S / "spike_syf1.docx")
for bbox, n, rows in report:
    print(f"--- blok {bbox}: tesseract {n} satır -> {len(rows)} satır")
    for r in rows:
        print("   ", r)
